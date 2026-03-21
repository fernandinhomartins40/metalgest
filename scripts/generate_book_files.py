from __future__ import annotations

import argparse
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Literal
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "doce habito de ser util.md"
DEFAULT_OUTPUT_DIR = ROOT / "uploads" / "book_exports"
DEFAULT_FONT_NAME = "Times New Roman"
DEFAULT_TITLE = "Livro"


@dataclass
class Layout:
    trim_w_cm: float
    trim_h_cm: float
    margin_cm: float
    body_font_size: float
    body_leading: float
    title_font_size: float
    part_font_size: float
    chapter_font_size: float
    interlude_font_size: float


@dataclass
class Block:
    kind: Literal["title", "part", "chapter", "interlude", "prologue", "epilogue", "paragraph"]
    text: str


def normalize_marker(text: str) -> str:
    folded = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    folded = folded.upper()
    folded = re.sub(r"\s+", " ", folded).strip()
    return folded


def sanitize_basename(name: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "_", name.strip())
    slug = re.sub(r"_+", "_", slug).strip("_")
    return slug or "livro"


def read_source(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def classify_heading(line: str, first_nonempty: bool) -> str | None:
    marker = normalize_marker(line)
    if first_nonempty:
        return "title"
    if marker.startswith("PARTE "):
        return "part"
    if marker.startswith("PART "):
        return "part"
    if marker.startswith("CAPITULO "):
        return "chapter"
    if marker.startswith("CHAPTER "):
        return "chapter"
    if marker.startswith("INTERLUDIO "):
        return "interlude"
    if marker.startswith("INTERLUDE "):
        return "interlude"
    if marker.startswith("PROLOGO"):
        return "prologue"
    if marker.startswith("PROLOGUE"):
        return "prologue"
    if marker.startswith("EPILOGO"):
        return "epilogue"
    if marker.startswith("EPILOGUE"):
        return "epilogue"
    return None


def parse_blocks(text: str) -> list[Block]:
    lines = [line.rstrip() for line in text.split("\n")]
    blocks: list[Block] = []
    paragraph_buffer: list[str] = []
    first_nonempty = True

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        paragraph = " ".join(part.strip() for part in paragraph_buffer if part.strip())
        paragraph = re.sub(r"\s{2,}", " ", paragraph).strip()
        if paragraph:
            blocks.append(Block("paragraph", paragraph))
        paragraph_buffer = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            continue
        heading_kind = classify_heading(stripped, first_nonempty)
        if heading_kind is not None:
            flush_paragraph()
            blocks.append(Block(heading_kind, stripped))
            first_nonempty = False
            continue
        paragraph_buffer.append(stripped)
        first_nonempty = False

    flush_paragraph()
    return blocks


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def set_docx_font(style, font_name: str, size: float, bold: bool = False, italic: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def setup_docx_styles(doc: Document, layout: Layout, font_name: str) -> None:
    normal = doc.styles["Normal"]
    set_docx_font(normal, font_name, layout.body_font_size)
    pf = normal.paragraph_format
    pf.first_line_indent = Cm(0.6)
    pf.line_spacing = Pt(layout.body_leading)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    set_docx_font(doc.styles["Title"], font_name, layout.title_font_size, bold=True)
    set_docx_font(doc.styles["Heading 1"], font_name, layout.part_font_size, bold=True)
    set_docx_font(doc.styles["Heading 2"], font_name, layout.chapter_font_size, bold=True)
    set_docx_font(doc.styles["Heading 3"], font_name, layout.interlude_font_size, bold=True, italic=True)


def summary_item_font_size(layout: Layout) -> float:
    return max(10.4, layout.body_font_size - 2.3)


def summary_item_leading(layout: Layout) -> float:
    return max(12.4, layout.body_leading - 5.0)


def summary_title_font_size(layout: Layout) -> float:
    return max(13.0, layout.chapter_font_size - 1.8)


def build_docx(blocks: list[Block], output_path: Path, layout: Layout, font_name: str, summary_title: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(layout.trim_w_cm)
    section.page_height = Cm(layout.trim_h_cm)
    section.top_margin = Cm(layout.margin_cm)
    section.bottom_margin = Cm(layout.margin_cm)
    section.left_margin = Cm(layout.margin_cm)
    section.right_margin = Cm(layout.margin_cm)

    setup_docx_styles(doc, layout, font_name)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    headings_for_summary = [
        block.text for block in blocks if block.kind in {"part", "prologue", "chapter", "interlude", "epilogue"}
    ]

    title_block = next((block for block in blocks if block.kind == "title"), None)
    if title_block is not None:
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(160)
        p.add_run(title_block.text)
        doc.add_page_break()

    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    title_run = p.add_run(summary_title)
    title_run.font.size = Pt(summary_title_font_size(layout))
    for item in headings_for_summary:
        row = doc.add_paragraph(style="Normal")
        row.paragraph_format.first_line_indent = Cm(0)
        row.paragraph_format.left_indent = Cm(0)
        row.paragraph_format.right_indent = Cm(0)
        row.paragraph_format.line_spacing = Pt(summary_item_leading(layout))
        row.paragraph_format.space_after = Pt(1.5)
        row.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = row.add_run(item)
        run.font.size = Pt(summary_item_font_size(layout))
    doc.add_page_break()

    first_content = True
    for block in blocks:
        if block.kind == "title":
            continue
        if block.kind in {"part", "prologue", "epilogue"}:
            if not first_content:
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(72)
            p.add_run(block.text)
            first_content = False
        elif block.kind == "chapter":
            doc.add_page_break()
            p = doc.add_paragraph(style="Heading 2")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.add_run(block.text)
        elif block.kind == "interlude":
            doc.add_page_break()
            p = doc.add_paragraph(style="Heading 3")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.add_run(block.text)
        else:
            p = doc.add_paragraph(style="Normal")
            if block.text.startswith("—"):
                p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(block.text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def register_pdf_font(font_name: str) -> tuple[str, str, str]:
    fonts_dir = Path(r"C:\Windows\Fonts")
    regular = fonts_dir / "times.ttf"
    bold = fonts_dir / "timesbd.ttf"
    italic = fonts_dir / "timesi.ttf"
    bold_italic = fonts_dir / "timesbi.ttf"
    if all(path.exists() for path in (regular, bold, italic, bold_italic)):
        pdfmetrics.registerFont(TTFont(font_name, str(regular)))
        pdfmetrics.registerFont(TTFont(f"{font_name}-Bold", str(bold)))
        pdfmetrics.registerFont(TTFont(f"{font_name}-Italic", str(italic)))
        pdfmetrics.registerFont(TTFont(f"{font_name}-BoldItalic", str(bold_italic)))
        return font_name, f"{font_name}-Bold", f"{font_name}-BoldItalic"
    return "Times-Roman", "Times-Bold", "Times-BoldItalic"


def build_pdf(
    blocks: list[Block],
    output_path: Path,
    layout: Layout,
    title: str,
    font_name: str,
    summary_title: str,
) -> None:
    regular_font, bold_font, bold_italic_font = register_pdf_font(font_name)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=(layout.trim_w_cm * cm, layout.trim_h_cm * cm),
        leftMargin=layout.margin_cm * cm,
        rightMargin=layout.margin_cm * cm,
        topMargin=layout.margin_cm * cm,
        bottomMargin=layout.margin_cm * cm,
        title=title,
        author="",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BookTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=layout.title_font_size,
        leading=layout.title_font_size + 5,
        alignment=TA_CENTER,
        spaceAfter=12,
        textColor=colors.black,
    )
    part_style = ParagraphStyle(
        "Part",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=layout.part_font_size,
        leading=layout.part_font_size + 4,
        alignment=TA_CENTER,
        spaceAfter=18,
        textColor=colors.black,
    )
    chapter_style = ParagraphStyle(
        "Chapter",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=layout.chapter_font_size,
        leading=layout.chapter_font_size + 4,
        alignment=TA_CENTER,
        spaceBefore=12,
        spaceAfter=12,
        textColor=colors.black,
    )
    interlude_style = ParagraphStyle(
        "Interlude",
        parent=styles["Heading2"],
        fontName=bold_italic_font,
        fontSize=layout.interlude_font_size,
        leading=layout.interlude_font_size + 4,
        alignment=TA_CENTER,
        spaceBefore=10,
        spaceAfter=10,
        textColor=colors.black,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=layout.body_font_size,
        leading=layout.body_leading,
        alignment=TA_JUSTIFY,
        firstLineIndent=0.6 * cm,
        spaceAfter=6,
        textColor=colors.black,
    )
    dialogue_style = ParagraphStyle(
        "Dialogue",
        parent=body_style,
        firstLineIndent=0,
    )
    summary_title_style = ParagraphStyle(
        "SummaryTitle",
        parent=part_style,
        fontSize=summary_title_font_size(layout),
        leading=summary_title_font_size(layout) + 3,
        spaceAfter=8,
    )
    summary_item_style = ParagraphStyle(
        "SummaryItem",
        parent=body_style,
        fontSize=summary_item_font_size(layout),
        leading=summary_item_leading(layout),
        alignment=TA_LEFT,
        firstLineIndent=0,
        leftIndent=0,
        rightIndent=0,
        spaceAfter=1.5,
    )

    story = []
    headings_for_summary = [
        block.text for block in blocks if block.kind in {"part", "prologue", "chapter", "interlude", "epilogue"}
    ]

    title_block = next((block for block in blocks if block.kind == "title"), None)
    book_title = title_block.text if title_block is not None else title or DEFAULT_TITLE
    story.append(Spacer(1, 4.8 * cm))
    story.append(Paragraph(escape(book_title), title_style))
    story.append(PageBreak())
    story.append(Paragraph(escape(summary_title), summary_title_style))
    for item in headings_for_summary:
        story.append(Paragraph(escape(item), summary_item_style))
    story.append(PageBreak())

    first_content = True
    for block in blocks:
        if block.kind == "title":
            continue
        text = escape(block.text)
        if block.kind in {"part", "prologue", "epilogue"}:
            if not first_content:
                story.append(PageBreak())
            story.append(Spacer(1, 2.3 * cm))
            story.append(Paragraph(text, part_style))
            first_content = False
        elif block.kind == "chapter":
            story.append(PageBreak())
            story.append(Paragraph(text, chapter_style))
        elif block.kind == "interlude":
            story.append(PageBreak())
            story.append(Paragraph(text, interlude_style))
        else:
            style = dialogue_style if block.text.startswith("—") else body_style
            story.append(Paragraph(text, style))

    def draw_page_number(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont(regular_font, 9)
        canvas.drawCentredString(doc_obj.pagesize[0] / 2, 0.85 * cm, str(canvas.getPageNumber()))
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)


def resolve_output_paths(source: Path, output_dir: Path, basename: str | None) -> tuple[Path, Path]:
    base = sanitize_basename(basename or source.stem)
    return output_dir / f"{base}.docx", output_dir / f"{base}.pdf"


def build_layout(args: argparse.Namespace) -> Layout:
    return Layout(
        trim_w_cm=args.trim_w_cm,
        trim_h_cm=args.trim_h_cm,
        margin_cm=args.margin_cm,
        body_font_size=args.body_font_size,
        body_leading=args.body_leading,
        title_font_size=args.title_font_size,
        part_font_size=args.part_font_size,
        chapter_font_size=args.chapter_font_size,
        interlude_font_size=args.interlude_font_size,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate DOCX and PDF book files from a markdown manuscript.")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE), help="Path to the markdown manuscript.")
    parser.add_argument("--title", default=DEFAULT_TITLE, help="Book title for metadata.")
    parser.add_argument("--basename", default=None, help="Base name for output files without extension.")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Directory for generated files.")
    parser.add_argument("--trim-w-cm", type=float, default=14.0, help="Trim width in centimeters.")
    parser.add_argument("--trim-h-cm", type=float, default=21.0, help="Trim height in centimeters.")
    parser.add_argument("--margin-cm", type=float, default=1.8, help="Page margin in centimeters.")
    parser.add_argument("--body-font-size", type=float, default=11.0, help="Body font size in points.")
    parser.add_argument("--body-leading", type=float, default=15.2, help="Body leading in points.")
    parser.add_argument("--title-font-size", type=float, default=24.0, help="Title page font size in points.")
    parser.add_argument("--part-font-size", type=float, default=16.0, help="Part/prologue/epilogue font size.")
    parser.add_argument("--chapter-font-size", type=float, default=15.0, help="Chapter font size.")
    parser.add_argument("--interlude-font-size", type=float, default=12.5, help="Interlude font size.")
    parser.add_argument("--font-name", default=DEFAULT_FONT_NAME, help="Preferred font family name.")
    parser.add_argument("--summary-title", default="Sumário", help="Summary page title.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source = Path(args.source).resolve()
    output_dir = Path(args.output_dir).resolve()
    docx_out, pdf_out = resolve_output_paths(source, output_dir, args.basename)
    layout = build_layout(args)

    text = read_source(source)
    blocks = parse_blocks(text)
    if not blocks:
        raise ValueError(f"No content parsed from {source}")

    build_docx(blocks, docx_out, layout, args.font_name, args.summary_title)
    build_pdf(blocks, pdf_out, layout, args.title, args.font_name, args.summary_title)

    print(f"DOCX: {docx_out}")
    print(f"PDF: {pdf_out}")


if __name__ == "__main__":
    main()
